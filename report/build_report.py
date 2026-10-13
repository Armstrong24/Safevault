from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2.2)

st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.5
for name, size in [("Heading 1", 16), ("Heading 2", 13)]:
    h = doc.styles[name]; h.font.name = "Times New Roman"; h.font.size = Pt(size); h.font.bold = True
    h.font.color.rgb = RGBColor(0, 0, 0)
    h.element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman"); h.element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(6)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)

def para(text, bold=False, align=None, size=None, italic=False, after=None):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align is not None: p.alignment = align
    else: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def rich(parts):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for t, b in parts:
        r = p.add_run(t); r.bold = b
    return p

def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.line_spacing = 1.3; p.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True; p.add_run(it[1])
        else: p.add_run(it)

def numbered(items):
    for it in items:
        p = doc.add_paragraph(style="List Number"); p.paragraph_format.line_spacing = 1.3; p.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True; p.add_run(it[1])
        else: p.add_run(it)

def table(headers, rows, widths, font=10.5, header_fill="D9D9D9"):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0]; repeat_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]; c.text = ""; p = c.paragraphs[0]; r = p.add_run(h); r.bold = True; r.font.size = Pt(font)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER; shade(c, header_fill)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            c = cells[i]; c.text = ""; p = c.paragraphs[0]; p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_after = Pt(0); r = p.add_run(str(v)); r.font.size = Pt(font)
            if i == 0: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2 == 1: pass
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def caption(text):
    p = para(text, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)

# ---------------- Title page ----------------
for _ in range(5): doc.add_paragraph()
para("SafeVault: Automated, Ransomware-Resilient Backup of Institutional Data to Cloud Storage",
     bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=18)
doc.add_paragraph()
para("Cloud Computing Project Report", align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
for _ in range(3): doc.add_paragraph()
para("Submitted by", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
t = doc.add_table(rows=5, cols=4); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
hdrs = ["Sr. No.", "Name", "Roll No.", "PRN"]
for i in range(5):
    for j in range(4):
        c = t.rows[i].cells[j]; c.text = ""; p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        if i == 0:
            r = p.add_run(hdrs[j]); r.bold = True
        elif j == 0:
            p.add_run(str(i))
        c.width = Cm([1.8, 6, 3.5, 4.2][j])
    t.rows[i].height = Cm(0.9)
doc.add_page_break()

# ---------------- Contents ----------------
doc.add_heading("Index", level=1)
toc = [("1", "Abstract"), ("2", "Introduction"), ("3", "Literature Review"), ("4", "Problem Statement"),
       ("5", "Project Objectives and Scope"), ("6", "System Diagram / Proposed Architecture"),
       ("7", "Project Module Explanation"), ("8", "System Requirements"), ("9", "Conclusion"),
       ("10", "Future Scope"), ("11", "References")]
table(["Sr. No.", "Topic", "Page No."], [(a, b, "") for a, b in toc], [2.2, 11, 2.8], font=12)
doc.add_page_break()

# ---------------- 1 Abstract ----------------
doc.add_heading("1. Abstract", level=1)
para("Colleges store a lot of important data like student marks, attendance, staff details and fee records. If this data gets deleted or locked by ransomware, it is very hard to get it back. In many departments, backups are still taken manually on a pen drive or a shared folder. These backups are on the same network as the original data, so if ransomware reaches the computer, it can encrypt or delete the backup as well.")
para("In this project we built SafeVault, a backup system that stores institutional data safely on Amazon Web Services (AWS). A Python program running on Ubuntu (WSL) takes a copy of the data, creates a SHA-256 hash for every file, compresses it and uploads it to a private Amazon S3 bucket in the Mumbai region. This runs automatically using cron. On the S3 bucket we turned on Versioning, Object Lock (Governance mode, 7 days), encryption, HTTPS-only access and Block Public Access. We also created two separate IAM users. The uploader user can only upload backups and cannot read, delete or unlock them. The recovery user can only read backups and update the dashboard.")
para("To check that recovery really works, we restore a selected clean backup into a new folder and compare the file hashes with a saved original copy. In our tests, all 5 out of 5 files matched. We also run a safe ransomware drill, where we encrypt a test copy of the data and check that the uploader key is not allowed to delete or read the backups. All results are shown on a live dashboard on Netlify, which also has a chatbot. We tested the system with 53 automated tests and with real runs on AWS. We used only sample (synthetic) data, and the system is ransomware-resilient for the cases we tested, not completely ransomware-proof.")
rich([("Keywords: ", True), ("Cloud backup, Ransomware, Amazon S3, Object Lock, Versioning, IAM, SHA-256, Data recovery", False)])

# ---------------- 2 Introduction ----------------
doc.add_heading("2. Introduction", level=1)
doc.add_heading("2.1 Background", level=2)
para("Every college department creates data daily, such as admission details, internal marks, attendance, timetables and fee receipts. If this data is lost even for a few days, it affects exams, results and audits. Taking backups is the usual solution, but a backup is useful only if it is safe when the original data is damaged.")
para("Ransomware is a type of malware that encrypts files and asks for money to unlock them. Many ransomware attacks also search for backup folders and connected drives and encrypt or delete them too, so the victim has no other option left. Educational institutions are often targeted because they use older systems and do not always have a dedicated security team.")
doc.add_heading("2.2 Why cloud computing", level=2)
para("Cloud storage like Amazon S3 keeps data away from the college network, so a problem on a local computer does not directly affect it. It is pay-as-you-use and there is no hardware to maintain. S3 also has useful security features. It can keep older versions of every file, it can lock files so they cannot be deleted for some days, and it lets us give each user only the permissions they need. Because of these features, we decided to use the cloud to build backups that can still be recovered even if the main computer is attacked.")
doc.add_heading("2.3 About our project", level=2)
para("SafeVault uses these cloud features for a college department. It takes backups automatically, checks them using hashes, stores them in a locked and versioned S3 bucket, and lets us restore and verify a clean copy. We also added a ransomware drill to test the system safely, and a dashboard that shows the results without showing any private data or keys. This project was done by our group of four students for the Cloud Computing subject.")

# ---------------- 3 Literature review ----------------
doc.add_heading("3. Literature Review", level=1)
para("We studied 16 recent research papers (published in 2025 and 2026) related to ransomware protection, backup and recovery, and secure cloud storage. The table below gives the paper title, the technique used and the gap we found in each paper.")
lit = [
 ("1", "Real-time Open-File Backup System with Machine-Learning Detection Model for Ransomware (Higuchi & Kobayashi, 2025)",
  "Backs up a file the moment it is opened, so a clean copy exists before ransomware can encrypt it",
  "Backups stay on the same machine; an attacker with admin rights can still delete them."),
 ("2", "Preventing Disruption of System Backup against Ransomware Attacks (Hou et al., 2025)",
  "Remembrall: watches for deletion of Windows shadow copies and uses it to detect ransomware",
  "Protects only local Windows shadow copies; no off-site or locked backup."),
 ("3", "Real-Time Detection and Recovery Method Against Ransomware Based on Simple Format Analysis (Kim, 2025)",
  "Checks file format and header to spot encrypted files and restores them in real time",
  "Recovery depends on the local system staying healthy; no cloud copy."),
 ("4", "A Machine Learning-Based Ransomware Detection Method for Attackers' Neutralization Techniques Using Format-Preserving Encryption (Lee et al., 2025)",
  "Machine learning models detect files encrypted with format-preserving encryption",
  "Only detects infected files; does not keep or restore a safe backup."),
 ("5", "SrFTL: Leveraging Storage Semantics for Effective Ransomware Defense in Flash-based SSDs (Zhu et al., 2025)",
  "Detection and file-level recovery built inside the SSD with a secure enclave",
  "Needs special SSD hardware; data is lost if the drive is stolen or fails."),
 ("6", "A Survey of Ransomware Detection Methods (Alzahrani et al., 2025)",
  "Reviews existing ransomware detection methods and their limits",
  "Focuses on detection; backup and verified recovery are not covered in depth."),
 ("7", "NAS Ransomware: Ransomware Targeting NAS Devices (Meurs et al., 2025)",
  "Studies 502 police-reported ransomware cases, 104 of them on network storage (NAS) devices",
  "Shows that connected backup drives get attacked, but gives no technical solution."),
 ("8", "A Proactive Camouflage-based Approach to Prevent Crypto-Ransomware: The Ransomware Prevention Toolkit (Alshaikh et al., 2026)",
  "Disguises important files as system files and runs a hidden monitor that stops ransomware",
  "Deliberately uses no backup storage, so data cannot be recovered if the trick fails."),
 ("9", "Enhancing Ransomware Resilience in Cloud-Based HR Systems through Moving Target Defense (Barach, 2026)",
  "Keeps changing containers, IP addresses and nodes in the cloud to confuse attackers",
  "Complex Kubernetes setup; does not focus on backup or restore checks."),
 ("10", "Modelling and Simulating Organizational Ransomware Recovery (Ilau et al., 2025)",
  "Simulation model to compare different ransomware recovery strategies in an organization",
  "Only a simulation; no working backup system is built or tested."),
 ("11", "From Backup Restoration to Minimum Viable Factory Recovery (Chiu, 2026)",
  "Literature review of ransomware recovery problems in manufacturing, e.g. untrusted restore points and no proof of recovery",
  "A framework only; needs real testing and practical tools."),
 ("12", "Enhancing Reliability, Security, and Scalability of Serverless Computing for Cloud Backup in Power Grids During Emergency Situations (Bilgaiyan et al., 2025)",
  "Studies serverless cloud computing as an emergency backup for power grids",
  "Focuses on emergency availability, not on ransomware or locked backups."),
 ("13", "Efficient Hash-Free Mobile Cloud Backup via Operation-Log Versioning (Pan et al., 2026)",
  "SolFS file system logs file changes so only changed data is uploaded to the cloud",
  "Aims at speed; has no ransomware protection or integrity checks."),
 ("14", "Enabling Flexible Multiple-Replica Cloud Storage Integrity Auditing for Cloud-based IoT Data (Chen & Hao, 2025)",
  "Checks that all copies of data in the cloud are intact, with low cost for the user",
  "Checks integrity only; does not stop deletion or test a restore."),
 ("15", "Smart Contract-based Public Integrity Auditing for Cloud Storage against Malicious Auditors (Tian et al., 2025)",
  "Uses blockchain smart contracts to make cloud integrity auditing transparent",
  "Heavy blockchain setup; no ransomware or backup recovery."),
 ("16", "Least-Privilege Role-Based Access Control Improvement for Cloud Container Security (Abdulraheem et al., 2026)",
  "Automatically designs least-privilege access roles for cloud container platforms",
  "Limited to Kubernetes roles; not applied to backup storage."),
]
table(["Paper ID", "Paper Title", "Technique Focus", "Identified Research Gap"], lit, [1.6, 5.6, 4.5, 4.5], font=10)
doc.add_heading("3.1 What we learned from the papers", level=2)
para("Most of the ransomware papers (1 to 9) focus on detecting the attack or keeping copies on the same computer or drive. If the attacker gets full control, or the drive fails, these copies can also be lost. Papers 7, 10 and 11 show that backups themselves are often attacked and that organizations need proof that recovery works. The cloud papers (12 to 16) improve cloud backup, integrity checking and access control, but each one handles only one part of the problem. None of them combine locked storage, limited permissions, hash checks and regular recovery testing in one simple system.")
para("In SafeVault we tried to cover this gap using normal AWS services. Our backups are stored off-site, are versioned and locked, the uploading computer cannot delete its own backups, every file has a hash, and we actually test the recovery instead of just assuming it works.")

# ---------------- 4 Problem statement ----------------
doc.add_heading("4. Problem Statement", level=1)
para("In many colleges, department data is backed up manually on pen drives or shared folders that are connected to the same computers that hold the original data. These backups are not taken regularly, are rarely tested, and can be encrypted or deleted by ransomware or by someone who gets access to the system. There is also no easy way to check whether a backup is complete, unchanged and can actually be restored.")
para("So, we need a simple and low-cost cloud backup system that:", after=2)
bullets([
 "takes backups automatically on a schedule,",
 "keeps backups safe even if the main computer and its keys are hacked,",
 "detects any change in the backed-up files using hashes,",
 "can restore a clean old version and check it against the original, and",
 "shows proof of protection and recovery without showing private data.",
])

# ---------------- 5 Objectives & scope ----------------
doc.add_heading("5. Project Objectives and Scope", level=1)
doc.add_heading("5.1 Objectives", level=2)
numbered([
 "To take backups automatically using cron, without any manual work.",
 "To create a SHA-256 hash for every file so that any change can be found.",
 "To store backups in a private S3 bucket with Versioning and Object Lock so they cannot be deleted easily.",
 "To use two separate IAM users with minimum permissions, so one stolen key cannot destroy the backups.",
 "To restore a selected clean backup and verify it using hashes.",
 "To safely test the system with a ransomware drill.",
 "To show the results on a live dashboard with a chatbot that answers questions about the project.",
])
doc.add_heading("5.2 Scope", level=2)
rich([("Included in the project: ", True), ("backup of file-based data (CSV files), Amazon S3 storage and its security settings, IAM users and policies, scheduled backups on Ubuntu (WSL), cloud restore and verification, ransomware drill on a test copy, and a public dashboard on Netlify.", False)])
rich([("Not included: ", True), ("live database backup, real student or personal data (we used sample data only), EC2 recovery server, CloudWatch and SNS alerts, and backup to multiple clouds. These are mentioned in Future Scope.", False)])
doc.add_heading("5.3 Limitations", level=2)
bullets([
 "Scheduled backups run only when the laptop and WSL are on and connected to the internet.",
 "In Governance mode, a special admin user can still remove the lock. Compliance mode is stricter but cannot be changed once set.",
 "We restore on the same laptop using a different user. In real use, a separate clean computer should be used.",
 "The system is ransomware-resilient for the cases we tested, but we do not claim it is 100% ransomware-proof.",
])

# ---------------- 6 Architecture ----------------
doc.add_heading("6. System Diagram / Proposed Architecture", level=1)
doc.add_picture("architecture.png", width=Cm(16))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption("Figure 6.1: Architecture of SafeVault")
para("Our system has four main parts:")
rich([("1. Institution side: ", True), ("The sample department data is stored on an Ubuntu (WSL) computer. Cron runs the backup program at a fixed time. The program takes a copy, creates hashes, compresses the files and uploads them using the uploader key.", False)])
rich([("2. AWS cloud: ", True), ("Backups are stored in a private S3 bucket in the Mumbai region (ap-south-1). Versioning, Object Lock, encryption, HTTPS-only access and Block Public Access are turned on. Each backup has an archive file and a completion record. IAM controls what each user is allowed to do.", False)])
rich([("3. Recovery: ", True), ("Using the recovery user, we download a selected clean backup, extract it into a new folder and compare the hashes with the original copy. The ransomware drill also uses this path.", False)])
rich([("4. Dashboard: ", True), ("After every command, the results (without any private details) are saved in a dashboard-data.json file in a separate bucket. The Netlify dashboard reads this file every 60 seconds. The chatbot uses a Netlify Function to call the Groq API, so the API key stays hidden.", False)])
doc.add_heading("6.1 Working steps", level=2)
numbered([
 "Cron starts the backup command.",
 "The program copies the data folder and creates a SHA-256 hash for each file.",
 "The files are compressed and uploaded to S3 as a new version.",
 "After the upload is successful, a completion record is uploaded.",
 "S3 locks the new versions for 7 days.",
 "The results are published to the dashboard.",
 "If something goes wrong, the recovery user restores a clean version and checks it.",
])

# ---------------- 7 Modules ----------------
doc.add_heading("7. Project Module Explanation", level=1)
mods = [
 ("7.1 Data Source Module", "This module creates sample department data such as student, staff and fee CSV files. It also saves the hash of each original file separately, which we use later to check if the restored data is correct. Since we use sample data, no real personal information is uploaded."),
 ("7.2 Scheduler Module", "This module sets up a cron job that runs the backup automatically, for example every 5 minutes. It makes sure two backups do not run at the same time. Every run is logged with its start time, end time and result. If a run fails, the error is saved and the last successful backup is still shown."),
 ("7.3 Backup Module", "This is the main program (safevault.py). It checks the settings, copies the data folder, creates a list of files with their SHA-256 hashes, compresses everything and uploads it to S3 using the AWS CLI. A completion record is uploaded only after the main file is uploaded successfully, so an incomplete backup is never treated as a valid one."),
 ("7.4 Cloud Storage and Security Module", "The S3 backup bucket is private. Versioning keeps every old version of a file. Object Lock (Governance mode, 7 days) stops locked versions from being deleted. Encryption protects the data at rest, and a bucket policy allows only HTTPS requests. Block Public Access is also on. The verify-protection command reads these settings from AWS and shows them, so we can confirm they are actually turned on."),
 ("7.5 IAM (Access Control) Module", "We created two IAM users. The first one, safevault-lite-uploader, can only upload files. It cannot read, list, delete or unlock anything. The second one, safevault-lite-recovery, can read backups and update only the dashboard file. So even if the uploader key is stolen from the laptop, the attacker cannot delete the old backups."),
 ("7.6 Recovery Module", "Using the recovery user, this module selects a clean backup version. We do not always pick the latest one, because the latest backup might already have damaged data. It downloads the backup, checks it, extracts it into a new folder and compares each file's hash with the original. It also checks that the CSV files open properly. In our tests, 5 out of 5 files matched."),
 ("7.7 Ransomware Drill Module", "This module safely tests a ransomware attack. It copies the data into a test folder and encrypts that copy. Then it tries to read, list, delete and unlock the backups using the uploader key, and AWS denies every request. After that, it restores the clean backup and verifies it. No real malware is used and the original data is not touched."),
 ("7.8 Dashboard Publisher Module", "This module collects the results of backups, protection checks, drills and restores, and saves them in a dashboard-data.json file. Before uploading, it removes private details like keys, account IDs, version IDs, bucket names and file paths. The file is uploaded to a separate bucket where only this one file is public."),
 ("7.9 Dashboard and Chatbot Module", "This is a web page hosted on Netlify. It reads the JSON file every 60 seconds and shows the current status, past runs, security settings and drill results. It also has a chatbot that answers questions about the project. The chatbot uses a Netlify Function to call the Groq API, and the key is stored in Netlify settings. If the API does not work, the chatbot gives some built-in answers."),
 ("7.10 Testing Module", "We wrote 53 automated tests. They include a full flow using a fake AWS CLI that follows the same permission rules as our real IAM users. After that, we ran real backups, IAM denial tests, Object Lock tests and cloud restores on our own AWS account."),
]
for h, body in mods:
    doc.add_heading(h, level=2); para(body)
doc.add_heading("7.11 Main commands", level=2)
table(["Command", "Module", "Output"], [
 ("lite.py schedule --every 5 --install", "Scheduler", "Cron job and run log"),
 ("lite.py run-backup", "Backup", "Backup file and completion record in S3"),
 ("lite.py verify-protection", "Storage and security", "S3 settings read from AWS"),
 ("lite.py cloud-restore", "Recovery", "Restored folder, 5/5 hashes match"),
 ("lite.py ransomware-drill", "Ransomware drill", "Access denied results and recovery"),
 ("lite.py publish / status", "Dashboard publisher", "Dashboard JSON and summary"),
], [5.8, 3.8, 6.4], font=10.5)

# ---------------- 8 Requirements ----------------
doc.add_heading("8. System Requirements", level=1)
doc.add_heading("8.1 Hardware requirements", level=2)
table(["Component", "Requirement"], [
 ("Processor", "Intel or AMD dual-core (64-bit)"),
 ("RAM", "Minimum 4 GB (8 GB recommended)"),
 ("Storage", "At least 2 GB free space"),
 ("Internet", "Stable internet connection"),
], [5, 11], font=11)
doc.add_heading("8.2 Software requirements", level=2)
table(["Software", "Use"], [
 ("Windows 10/11 with Ubuntu (WSL)", "To run the backup program and cron"),
 ("Python 3", "Main programming language"),
 ("AWS CLI v2", "To upload, download and check S3 settings"),
 ("Cron", "To schedule automatic backups"),
 ("HTML, CSS, JavaScript", "To build the dashboard"),
 ("Netlify Functions (Node.js)", "For the chatbot backend"),
 ("Python unittest", "For testing"),
 ("Web browser", "To open the dashboard"),
], [6.5, 9.5], font=11)
doc.add_heading("8.3 Cloud services", level=2)
table(["Service", "Use in project"], [
 ("Amazon S3 (backup bucket)", "Stores backups with Versioning, Object Lock, encryption and no public access"),
 ("Amazon S3 (dashboard bucket)", "Stores one public JSON file for the dashboard"),
 ("AWS IAM", "Uploader and recovery users with limited permissions"),
 ("AWS Region ap-south-1 (Mumbai)", "Region used for all buckets"),
 ("Netlify", "Hosting the dashboard and the chatbot function"),
 ("Groq API", "AI chatbot answers"),
], [6.5, 9.5], font=11)

# ---------------- 9 Conclusion ----------------
doc.add_heading("9. Conclusion", level=1)
para("In this project we built SafeVault, a cloud backup system that protects college data from ransomware and accidental deletion. We used AWS S3 features like Versioning, Object Lock and encryption along with separate IAM users. Because of this, the computer that uploads the backups cannot read, delete or unlock them, which is the main way ransomware destroys backups.")
para("We did not just store the backups, we also checked them. Every backup has SHA-256 hashes, every restore is compared with the original files, and the ransomware drill showed that the attacker's key is blocked and the clean data can be restored. We tested this with 53 automated tests and real runs on AWS, and the results are shown on a live dashboard. Through this project we learned how cloud services can be used to make a simple, low-cost and reliable backup system for institutions.")

# ---------------- 10 Future scope ----------------
doc.add_heading("10. Future Scope", level=1)
bullets([
 "Run backups using AWS Lambda or an EC2 server, so they work even when the laptop is off.",
 "Add CloudWatch and SNS to send email or SMS alerts when a backup fails.",
 "Use Compliance mode in Object Lock for very important data, so even admins cannot remove the lock.",
 "Copy backups to another AWS region or account for extra safety (3-2-1 backup rule).",
 "Move old backups to S3 Glacier to reduce cost.",
 "Add support for database backups like MySQL.",
 "Detect suspicious changes in files automatically before they are backed up.",
 "Restore data on a separate clean virtual machine and add multi-factor authentication for recovery.",
 "Extend the system to multiple departments with separate dashboards.",
])

# ---------------- 11 References ----------------
doc.add_heading("11. References", level=1)
para("References are written in APA (7th edition) style and listed in alphabetical order.")
# each reference: list of (text, italic) parts
refs = [
 [("Abdulraheem, W. K., Ibbini, E. M., Kanaker, H., Smadi, S., Karim, N. A., Fakhouri, H. N., Albdour, L., & Fakhouri, S. (2026). Least-privilege role-based access control improvement for cloud container security. ",0),("Computers, 15",1),("(5), Article 326. https://doi.org/10.3390/computers15050326",0)],
 [("Alshaikh, H., Hefny, H., & Ramadan, N. (2026). A proactive camouflage-based approach to prevent crypto-ransomware: The ransomware prevention toolkit (RPTK). ",0),("International Journal of Information Security, 25",1),("(4), Article 135. https://doi.org/10.1007/s10207-026-01296-0",0)],
 [("Alzahrani, S., Xiao, Y., Asiri, S., Zheng, J., & Li, T. (2025). A survey of ransomware detection methods. ",0),("IEEE Access, 13",1),(", 57943–57982. https://doi.org/10.1109/ACCESS.2025.3556187",0)],
 [("Amazon Web Services. (n.d.). ",0),("Locking objects with Object Lock",1),(". Amazon S3 User Guide. https://docs.aws.amazon.com/AmazonS3/latest/userguide/object-lock.html",0)],
 [("Amazon Web Services. (n.d.). ",0),("Retaining multiple versions of objects with S3 Versioning",1),(". Amazon S3 User Guide. https://docs.aws.amazon.com/AmazonS3/latest/userguide/Versioning.html",0)],
 [("Amazon Web Services. (n.d.). ",0),("Security best practices in IAM",1),(". AWS Identity and Access Management User Guide. https://docs.aws.amazon.com/IAM/latest/UserGuide/best-practices.html",0)],
 [("Barach, J. (2026). Enhancing ransomware resilience in cloud-based HR systems through moving target defense. ",0),("Computers, Materials & Continua, 86",1),("(2), 1–23. https://doi.org/10.32604/cmc.2025.071705",0)],
 [("Bilgaiyan, S., Aryan, A., Nayak, N. K., & Mishra, P. (2025). Enhancing reliability, security, and scalability of serverless computing for cloud backup in power grids during emergency situations. ",0),("SN Computer Science, 6",1),("(4), Article 390. https://doi.org/10.1007/s42979-025-03941-2",0)],
 [("Chen, G., & Hao, R. (2025). Enabling flexible multiple-replica cloud storage integrity auditing for cloud-based IoT data. ",0),("Journal of Systems Architecture, 163",1),(", Article 103404. https://doi.org/10.1016/j.sysarc.2025.103404",0)],
 [("Chiu, C. Y. (2026). From backup restoration to minimum viable factory recovery: A systematization of ransomware recovery in manufacturing systems. ",0),("International Journal of Critical Infrastructure Protection, 54",1),(", Article 100882. https://doi.org/10.1016/j.ijcip.2026.100882",0)],
 [("Higuchi, K., & Kobayashi, R. (2025). Real-time open-file backup system with machine-learning detection model for ransomware. ",0),("International Journal of Information Security, 24",1),("(1), Article 54. https://doi.org/10.1007/s10207-024-00966-1",0)],
 [("Hou, Y., Guo, L., Zhou, C., Zhang, Q., Liu, W., Sun, C., & Jiang, Y. (2025). Preventing disruption of system backup against ransomware attacks. ",0),("Proceedings of the ACM on Software Engineering, 2",1),("(ISSTA), 229–249. https://doi.org/10.1145/3728880",0)],
 [("Ilau, M.-C., Baldwin, A., Caulfield, T., & Pym, D. (2025). Modelling and simulating organizational ransomware recovery: Structure, methodology, and decisions. ",0),("Journal of Cybersecurity, 11",1),("(1), Article tyaf035. https://doi.org/10.1093/cybsec/tyaf035",0)],
 [("Kim, J. (2025). Real-time detection and recovery method against ransomware based on simple format analysis. ",0),("Information, 16",1),("(9), Article 739. https://doi.org/10.3390/info16090739",0)],
 [("Lee, J., Kim, J., Jeong, H., & Lee, K. (2025). A machine learning-based ransomware detection method for attackers' neutralization techniques using format-preserving encryption. ",0),("Sensors, 25",1),("(8), Article 2406. https://doi.org/10.3390/s25082406",0)],
 [("Meurs, T., Hoheisel, R., Junger, M., & Abhishta, A. (2025). NAS ransomware: Ransomware targeting NAS devices. ",0),("Global Crime, 27",1),("(2–3), 135–158. https://doi.org/10.1080/17440572.2025.2591628",0)],
 [("Netlify. (n.d.). ",0),("Functions overview",1),(". Netlify Docs. https://docs.netlify.com/functions/overview/",0)],
 [("Pan, R., Liang, Y., Li, L., Du, H., Kuo, T.-W., & Xue, C. J. (2026). Efficient hash-free mobile cloud backup via operation-log versioning. ",0),("ACM Transactions on Storage, 22",1),("(3), 1–26. https://doi.org/10.1145/3789213",0)],
 [("Tian, H., Gan, N., Peng, F., Quan, H., Chang, C.-C., & Vasilakos, A. V. (2025). Smart contract-based public integrity auditing for cloud storage against malicious auditors. ",0),("Future Generation Computer Systems, 166",1),(", Article 107709. https://doi.org/10.1016/j.future.2025.107709",0)],
 [("Zhu, W., Hernandez, G., Garcia, W., Tian, D. J., Rampazzi, S., & Butler, K. R. B. (2025). SrFTL: Leveraging storage semantics for effective ransomware defense in flash-based SSDs. ",0),("ACM Transactions on Storage, 21",1),("(4), 1–42. https://doi.org/10.1145/3767322",0)],
]
for parts in refs:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5; p.paragraph_format.left_indent = Cm(1.27); p.paragraph_format.first_line_indent = Cm(-1.27)
    p.paragraph_format.space_after = Pt(6)
    for t, it in parts:
        r = p.add_run(t); r.font.size = Pt(11); r.italic = bool(it)

# page numbers in footer
for s in doc.sections:
    fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for tag, txt in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); run._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; run._r.append(e)

cp = doc.core_properties; cp.author = "Arnav Kale"; cp.last_modified_by = "Arnav Kale"; cp.title = "SafeVault Project Report"; cp.comments = ""; cp.keywords = ""
doc.save("SafeVault_Project_Report.docx")
print("saved")
